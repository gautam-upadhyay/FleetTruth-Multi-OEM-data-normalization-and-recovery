"""Populate the retained Motorq solution template without changing its source package."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
import math
import pypdfium2 as pdfium
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, white

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT.parent / 'Motorq_Hackathon_Solution_Document_Template.docx'
EXPECTED_HASH = '8ebd4d17b3560894875ea19c4fecfa55ac66b4373e0e107db25cf651d35885a2'
OUT = ROOT / 'artifacts/submission'
TMP = ROOT / 'tmp/pdfs'
ASSETS = TMP / 'figures'
for folder in [OUT, TMP, ASSETS]:
    folder.mkdir(parents=True, exist_ok=True)
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == EXPECTED_HASH

source = Document(REFERENCE)
doc = Document(REFERENCE)
paras = [deepcopy(p._p) for p in source.paragraphs]
tables = [deepcopy(t._tbl) for t in source.tables]
headings = {p.text: deepcopy(p._p) for p in source.paragraphs if p.style.name.startswith('Heading')}
body = doc._element.body
for element in list(body):
    if element.tag != qn('w:sectPr'):
        body.remove(element)


def insert(element):
    body.insert(len(body) - 1, element)


def text(value, bold=False, size=None, keep=False):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.widow_control = True
    run = p.add_run(value)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return p


def heading(name, page=False):
    p = Paragraph(deepcopy(headings[name]), body)
    p.paragraph_format.page_break_before = False
    insert(p._p)


def sub(value):
    return text(value, bold=True, keep=True)


def code(value):
    p = text(value, size=9.5)
    p.paragraph_format.line_spacing = 1.0
    for r in p.runs:
        r.font.name = 'Consolas'
    return p


def matrix(headers, rows, widths=None, template=None, size=10):
    index = template if template is not None else {3: 3, 4: 0, 5: 8, 7: 2}[len(headers)]
    table = Table(deepcopy(tables[index]), body)
    insert(table._tbl)
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for _ in rows:
        table.add_row()
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    for ri, values in enumerate([headers, *rows]):
        row = table.rows[ri]
        prop = row._tr.get_or_add_trPr()
        prop.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            prop.append(OxmlElement('w:tblHeader'))
        for ci, value in enumerate(values):
            cell = row.cells[ci]
            if widths:
                cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.bold = ri == 0
                    run.italic = False
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def figure(path, caption, width=6.35):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width))
    for e in r._r.iter():
        if e.tag.endswith('}docPr'):
            e.set('descr', caption)
    cap = text(caption, size=9.5)
    cap.paragraph_format.space_after = Pt(10)


def diagram(name, nodes, edges, height=3.4):
    c = canvas.Canvas(str(ASSETS/(name+'.pdf')), pagesize=(460.8,height*72))
    sx,sy=4.608,height*.72
    points = {}
    for key, x, y, w, h, label in nodes:
        points[key] = (x, y, w, h)
        c.setFillColor(HexColor('#f2f5f5'));c.setStrokeColor(HexColor('#708280'));c.setLineWidth(.7)
        c.rect(x*sx,y*sy,w*sx,h*sy,fill=1,stroke=1)
        center_text(c,(x+w/2)*sx,(y+h/2)*sy,label,9.1)
    for start, end, label in edges:
        x,y,w,h = points[start]
        u,v,a,b = points[end]
        if abs((x+w/2)-(u+a/2)) > abs((y+h/2)-(v+b/2)):
            p1 = (x+w if u>x else x, y+h/2)
            p2 = (u if u>x else u+a, v+b/2)
            offset = (0, 3)
        else:
            p1 = (x+w/2, y if v<y else y+h)
            p2 = (u+a/2, v+b if v<y else v)
            offset = (2.4, 0)
        arrow(c,(p1[0]*sx,p1[1]*sy),(p2[0]*sx,p2[1]*sy))
        if label:
            center_text(c,((p1[0]+p2[0])/2+offset[0])*sx,((p1[1]+p2[1])/2+offset[1])*sy,label,7.7,True)
    c.save()
    return raster(name)


def center_text(c,x,y,value,size=9,background=False):
    lines=value.split('\n')
    c.setFont('Helvetica',size)
    if background:
        width=max(c.stringWidth(line,'Helvetica',size) for line in lines)
        c.setFillColor(white);c.rect(x-width/2-2,y-size*len(lines)/2-2,width+4,size*len(lines)+4,fill=1,stroke=0)
    c.setFillColor(HexColor('#172a28'))
    for i,line in enumerate(lines):
        c.drawCentredString(x,y+(len(lines)-1)*size*.6-i*size*1.2-size*.32,line)


def arrow(c,p1,p2):
    c.setStrokeColor(HexColor('#466763'));c.setFillColor(HexColor('#466763'));c.setLineWidth(.8)
    c.line(*p1,*p2)
    angle=math.atan2(p2[1]-p1[1],p2[0]-p1[0])
    p=c.beginPath();p.moveTo(*p2)
    p.lineTo(p2[0]-6*math.cos(angle-.42),p2[1]-6*math.sin(angle-.42))
    p.lineTo(p2[0]-6*math.cos(angle+.42),p2[1]-6*math.sin(angle+.42));p.close()
    c.drawPath(p,fill=1,stroke=0)


def raster(name):
    pdf=pdfium.PdfDocument(ASSETS/(name+'.pdf'))
    path=ASSETS/(name+'.png')
    pdf[0].render(scale=3).to_pil().save(path)
    pdf.close()
    return path


context = diagram('context', [
    ('users',2,72,36,24,'Integration engineer\nFleet operator / administrator'),
    ('system',55,72,43,24,'FleetTruth\nNormalization and recovery'),
    ('oem',2,9,36,26,'Four synthetic OEM streams\nReal OEM connectors proposed'),
    ('idp',55,9,43,26,'Organizational identity provider\nOIDC adapter unverified')],
    [('users','system','Browser / HTTP locally'),('oem','system','JSON telemetry'),('idp','system','OIDC / HTTPS target')],2.9)
containers = diagram('containers', [
    ('web',1,78,29,20,'React workspace\nnginx :3000'),('api',38,78,29,20,'Spring API / workers\nJWT and domain rules'),
    ('sim',1,42,29,20,'Synthetic simulator\n100K enrolled / 192 live'),('kafka',38,42,29,20,'Kafka raw topic\ntenant:VIN key'),
    ('sql',73,42,26,20,'PostgreSQL\npgvector'),('sinks',73,2,26,23,'Redis / optional\nCassandra projections'),
    ('batch',38,2,29,23,'Export / Parquet\nDuckDB or Spark')],
    [('web','api','HTTP / JSON'),('sim','kafka','Kafka protocol'),('kafka','api','Kafka consumer'),
     ('api','sql','JDBC'),('sql','sinks','Outbox worker'),('sql','batch','NDJSON export')],4.0)
er = diagram('relational-model', [
    ('tenant',1,78,27,19,'TENANTS'),('fleet',36,78,27,19,'FLEETS'),('vehicle',72,78,27,19,'VEHICLES'),
    ('oem',1,43,27,19,'OEMS'),('model',36,43,27,19,'VEHICLE_MODELS'),('state',72,43,27,19,'VEHICLE_STATE'),
    ('mapping',1,5,27,20,'MAPPINGS\nversion + approval'),('events',36,5,27,20,'RAW_EVENTS\nsource evidence'),
    ('rev',72,5,27,20,'REVISIONS\nALERTS / OUTBOX')],
    [('tenant','fleet','1:N'),('fleet','vehicle','1:N'),('oem','model','1:N'),('model','vehicle','1:N'),
     ('vehicle','state','1:1'),('mapping','events','interprets'),('events','rev','1:N')],3.6)
deployment = diagram('deployment', [
    ('host',1,72,98,24,'Verified local Docker network | loopback ports only\nWeb 3000 | API 8082 | PostgreSQL, Kafka and Redis internal'),
    ('monitor',1,38,98,20,'Optional observability profile\nPrometheus 9090 | Grafana 3001 | Jaeger 16686'),
    ('target',1,1,98,25,'Proposed Kubernetes deployment\nAPI 2-8 replicas | web 2 replicas | external data services\nIngress TLS, managed identity, vault and HA still to configure')],
    [('host','monitor','Metrics and OTLP'),('monitor','target','Portable images and Helm assets')],3.2)
layers = diagram('layers', [
    ('api',3,77,94,19,'Presentation | FleetController and SecurityConfig'),
    ('service',3,42,94,21,'Application | FleetService, ReplayWorker, PrivacyService'),
    ('domain',3,4,44,21,'Domain rules\nNormalizer / VinValidator'),
    ('infra',53,4,44,21,'Infrastructure adapters\nJDBC / Kafka / projections')],
    [('api','service','DTOs and authorized calls'),('service','domain','Validation / transforms'),('service','infra','Persistence / delivery')],2.8)


def sequence(name, labels, steps):
    c=canvas.Canvas(str(ASSETS/(name+'.pdf')),pagesize=(460.8,259.2))
    xs=[38+i*384/(len(labels)-1) for i in range(len(labels))]
    for x,label in zip(xs,labels):
        center_text(c,x,245,label,9.5)
        c.setStrokeColor(HexColor('#afbdbb'));c.setDash(3,3);c.line(x,14,x,230);c.setDash()
    for i,(a,b,label) in enumerate(steps):
        y=220-i*31
        arrow(c,(xs[a],y),(xs[b],y))
        center_text(c,230,y+9,label,8.3,True)
    c.save()
    return raster(name)


seq_ingest=sequence('sequence-ingest',['Producer','Kafka','Normalizer','SQL / outbox'],[
    (0,1,'1  Send event with stable event ID'),(1,2,'2  Consume tenant:VIN record'),
    (2,3,'3  Check ownership and unique tenant/event key'),(2,3,'4  Commit event, state, revision, alert and outbox'),
    (2,1,'5  Acknowledge after successful service return'),(1,2,'6  Crash/redelivery repeats the same event ID'),
    (3,2,'7  Duplicate detected; no second business effect')])
seq_replay=sequence('sequence-recovery',['OEM event','Quarantine','Engineer','Replay / SQL'],[
    (0,1,'1  New schema has no approved mapping'),(1,2,'2  Inspect retained payload and incident'),
    (2,3,'3  Preview candidate rules against samples'),(3,2,'4  Return validation and canonical comparison'),
    (2,3,'5  Explicitly approve then request replay'),(3,1,'6  Read retained events in checkpointed pages'),
    (3,2,'7  Report recovery, alert evidence and audit')])

benchmark=json.loads((ROOT/'evidence/local-api-benchmark.json').read_text())
metrics=['p50','p95','p99'];values=[benchmark[k+'Ms'] for k in metrics]
c=canvas.Canvas(str(ASSETS/'latency.pdf'),pagesize=(460.8,201.6))
center_text(c,230,190,'Client observed latency in milliseconds',10)
for i,(name,value,color) in enumerate(zip(metrics,values,['#91bab3','#477f77','#244f49'])):
    x=100+i*125
    c.setFillColor(HexColor(color));c.rect(x-28,28,56,value*.6,fill=1,stroke=0)
    center_text(c,x,28+value*.6+11,f'{value:.2f}',11)
    center_text(c,x,14,name,10)
c.save();raster('latency')

# Retain the original cover, linked outline and their page furniture.
cover_values={1:'FleetTruth Solution Document',2:'Submission Format: PDF',3:'To be Submitted by: ______________________________',
4:'Team Members & Roles: ___________________________',5:'Problem Space Chosen: Multi-OEM data normalization and recovery',
6:'Repository URL: __________________________________',7:'Demo Video URL (5 min maximum): __________________',8:'Date of Submission: 02/10/2026'}
for i in range(9):
    el=deepcopy(paras[i]);p=Paragraph(el,body)
    if i in cover_values:
        first=p.runs[0] if p.runs else p.add_run()
        first.text=cover_values[i]
        for run in p.runs[1:]:run.text=''
    insert(el)
doc.add_page_break()
for i in range(10,43):insert(deepcopy(paras[i]))
doc.add_page_break()

heading('1. Executive Summary')
text('FleetTruth helps fleet integration engineers keep vehicle signals trustworthy when manufacturers change telemetry formats. A battery fraction such as 0.07 must become 7%, and a value with unknown semantics must remain untrusted until an engineer approves its interpretation. Fleet operators need this distinction before acting on charging, vehicle health or operational alerts.')
text('We implemented a versioned normalization and recovery workflow that retains original events, quarantines incompatible schemas, validates mapping rules against source samples, requires human approval and replays retained events. The interface exposes source-to-canonical comparisons, recovery checkpoints, alert evidence and an audit trail. A read-only assistant retrieves workspace evidence; an advisory classifier helps prioritize drift review.')
text('The local registry contains 100,000 synthetic vehicles; the presentation stream rotates through 192 vehicles at approximately 12 events/sec. Clean backend verification passed 63 tests, including real PostgreSQL, Kafka, Redis, pgvector and Cassandra integration, Pact and Cucumber. Twenty browser tests passed in each of the native and Docker deployments. All ten selected core classes exceed 80% line and instruction coverage. A bounded 90-request benchmark measured p95 84.19 ms and p99 205.77 ms. Full-scale ingestion, HA and production security remain future validation work.')

heading('2. Problem Statement & Validation')
heading('2.1 Problem Statement')
text('An integration engineer supporting a mixed-manufacturer fleet needs a way to detect and safely repair changing telemetry contracts because a syntactically valid message can carry different units or field meanings. The immediate measurable impact is a backlog of unusable events and stale or misleading vehicle state. Financial loss and customer remediation time have not yet been measured.')
text('The primary user is the fleet integration engineer. Fleet operators consume trusted state and alerts; administrators manage tenant access, audit and privacy. The synthetic demonstration spans Aster Motors, Helix Automotive, Nord Electric and Vertex Mobility. These are invented manufacturers.')
text('The concrete failure is a Helix change from a percentage field to soc_fraction. A legacy parser may drop the new value, display 0.07% or retain old state. FleetTruth creates an incident, marks the affected state untrusted and retains the source message. After an engineer validates scale 100, replay interprets 0.07 as 7%. A moving low-battery vehicle can then produce an alert linked to its mapping version.')
heading('2.2 Evidence & Validation')
text('Validation uses synthetic simulation, executable workflow tests and local measurements. No customer interviews, production incident dataset or measured customer ROI were collected. Existing alternatives considered were separate OEM portals, manually maintained connector code and general fleet-data platforms. We have not benchmarked commercial products; differentiation here is the inspectable approval-and-replay workflow demonstrated by this implementation.')
matrix(['Evidence / Assumption','Source or Method','What It Shows','Confidence'],[
['Schema and unit changes can break interpretation','Injected Helix schema change; NormalizerTest and browser recovery journey','Quarantine, approved conversion and recovery are executable','High in tested cases'],
['Recovery preserves state ordering and duplicate safety','WorkflowTest / PostgresIntegrationTest','Late replay does not replace newer state; duplicate IDs do not duplicate effects','High in tested cases'],
['Read latency is suitable for a local demo','90 reads, concurrency 3; local-api-benchmark.json','p95 84.19 ms; p99 205.77 ms; 0 errors','High for this sample'],
['Engineers save time and avoid incorrect decisions','Proposed pilot comparison with prior workflow','Business value hypothesis; savings unmeasured','Unvalidated']],[1.55,1.65,2.0,1.25])
heading('2.3 Impact & Success Metrics')
matrix(['Metric','Baseline Today','Target','How Measured / Estimated'],[
['Incorrect acceptance of unknown schemas','Risk demonstrated with synthetic contract changes','Zero acceptance before approval in tested fixtures','Normalizer and workflow assertions'],
['Duplicate business effects','Potential under at-least-once delivery','Zero repeated effects for the same tenant/event ID','H2 and PostgreSQL tests'],
['API latency','Earlier local p95 569.33 ms; workload differs','Case-study p95 <200 ms; p99 <500 ms','Current bounded read sample; no controlled speedup claim'],
['Time to recover a contract incident','Customer baseline not measured','Measure reduction during pilot','Incident-to-replay completion time with event reconciliation']],[1.5,1.7,1.6,1.65])
text('At 10K and 100K vehicles, the mechanism is the same: a changed contract is repaired once per OEM/schema version, then applied to retained events. Exposure grows with the number of affected vehicles and messages. At an assumed one event per vehicle per second, 10K vehicles create 864 million events/day and 100K create 8.64 billion. These are capacity estimates, not achieved rates. Potential safety, charging and maintenance benefits require operator validation; no carbon reduction or financial saving is claimed.')

heading('3. Solution Description',page=True)
heading('3.1 Solution Overview & User Journey')
text('The workspace makes the state of a telemetry contract visible alongside the vehicle state it affects. The engineer can inspect what arrived, what conversion is proposed and which retained events will be recovered. The source remains available for audit after normalization.')
text('Journey: a synthetic event enters Kafka; the normalizer selects an approved exact-version mapping; an unknown schema enters quarantine; an engineer previews the mapping against retained samples; explicit approval creates the trusted revision; a durable replay job reprocesses the backlog; fleet alerts and audit records expose the recovered evidence.')
figure(ROOT/'artifacts/demo/frame-mapping.png','Figure 1. Working Mapping studio: source payload, canonical conversion and 20 retained validation samples. Screenshot from the supplied local recording.')
heading('3.2 Key Value Proposition')
text('The customer job is to maintain a usable multi-manufacturer fleet feed while OEM contracts evolve. FleetTruth reduces the diagnostic ambiguity between missing data, wrong units and stale state. It adds a recoverable workflow: retained events can be reinterpreted after approval without asking the vehicle to retransmit them. Original evidence and mapping versions make an alert explainable to an operator.')
text('The distinctive implementation combines explicit contract approval, side-by-side payload comparison and resumable recovery in one workspace. Human review remains necessary because successful sample validation cannot establish that a mapping agrees with a real OEM specification.')
heading('3.3 Innovative Ideas')
text('1. Evidence-preserving normalization: unknown formats remain inspectable in quarantine. The validation and replay journeys demonstrate how one approved rule recovers earlier events while preserving original timestamps and revision provenance.')
text('2. Recovery-aware trust: event-time and sequence ordering prevent an older replay from replacing newer vehicle state or alert evidence. Duplicate and replay tests exercise this behavior in H2 and PostgreSQL.')
text('3. Transactionally maintained read summaries: TOTAL, MINUTE and DAY counters change with ingestion, replay and local erasure. Ledger-reconciliation, rollback and concurrent-writer tests check correctness. This removes repeated dashboard history scans, while introducing a documented summary-row contention limit at large write rates. These are project design choices, not claims of research novelty.')

heading('4. Feature List',page=True)
text('Timestamps below are approximate positions in the 114.88-second supplied recording. Done refers to the implemented local scope. Production scale and deployment are assessed separately. Code aliases expand immediately below the table.')
matrix(['ID','Feature','User Story','Priority','Status','Code Path','Video'],[
['F01','Trust overview','Operator inspects quality and critical alerts','Must','Done','FS; UI','00:03'],
['F02','Profile / themes','User views identity and appearance','Should','Done','UP; TH','00:10'],
['F03','OEM contracts','Engineer finds changed schema','Must','Done','FS; DM','00:27'],
['F04','Mapping preview','Engineer compares units before approval','Must','Done','N; FS; UI','00:40'],
['F05','Human approval','Engineer authorizes a mapping version','Must','Done','FC; FS','00:56'],
['F06','Durable replay','Engineer recovers retained events','Must','Done','PW; FS','01:02'],
['F07','Alert evidence','Operator inspects normalized evidence','Must','Done','FS; FC; UI','01:10'],
['F08','Read-only assistant','User retrieves audited workspace facts','Should','Done','AS; KS','01:26'],
['F09','Daily analytics','Operator exports persisted summaries','Should','Done','TR; UI','01:41'],
['F10','Local erasure','Admin confirms a synthetic VIN purge','Should','Partial','PS; GOV','Not in clip'],
['F11','Production HA / OIDC','Organization deploys secure fleet service','Could','Planned','Helm / OIDC','Not in clip'],
['F12','Autonomous approvals','Agent changes fleet interpretation','Won\'t','Planned exclusion','No mutation tool','Not applicable']],[.35,.88,1.61,.55,.57,1.0,.69],template=2,size=9.5)
text('Java source root: api/src/main/java/io/fleettruth/. FC = api/FleetController.java; N = domain/Normalizer.java; FS = service/FleetService.java; PW = service/ReplayWorker.java; DM = service/DriftModel.java; AS = service/AssistantService.java; KS = service/KnowledgeService.java; TR = service/TelemetryRollups.java; PS = service/PrivacyService.java.',size=10)
text('Web source root: web/src/. UI = pages.tsx; UP = UserProfile.tsx; TH = theme.tsx; GOV = Governance.tsx. Deployment source: infra/helm/fleettruth and web/src/EnterpriseLogin.tsx.',size=10)

heading('5. Solution Architecture (High-Level Design)',page=True)
heading('5.1 Architecture Overview')
text('FleetTruth runs as a modular Spring application with adapters and scheduled workers. PostgreSQL holds authoritative decisions and retained event evidence; Kafka transports events; Redis and optional Cassandra receive projections through a durable outbox. This deployment keeps the demonstrator understandable while exposing boundaries for later worker separation.')
figure(context,'Figure 2. System context. Real OEM connectors and organizational identity are future integrations; the verified workflow uses synthetic streams and local JWT login.')
figure(containers,'Figure 3. Local containers and protocols. API reads use SQL. Redis is a projection; Cassandra is optional. pgvector holds authored runbooks. Native mode substitutes H2 and direct intake.')
text('One event travels from the simulator to Kafka, then through ownership/schema validation and a SQL transaction. The outbox later forwards normalized events and sink projections. The browser polls every two seconds. Per-hop latency has not been measured. The design target is under two seconds to the dashboard and under five seconds to a critical alert; polling interval alone does not prove either end-to-end target.')
heading('5.2 Technology Stack & Justification')
matrix(['Layer','Choice','Why This, and What You Rejected'],[
['Ingestion / messaging','Kafka 3.9.1; JSON; tenant:VIN key','Retained records and per-key ordering fit replay. A request-only pipeline cannot buffer unavailable consumers.'],
['Stream / batch processing','Java normalizer; Parquet + DuckDB; Spark script','Explicit rules remain auditable. Local batch is inexpensive to reproduce; Spark is reserved for cluster-scale history.'],
['Relational / NoSQL / cache / vector','PostgreSQL + pgvector; H2 locally; Redis; optional Cassandra','ACID control decisions, disposable latest-state projection, time-bucketed telemetry and tiny exact vector retrieval.'],
['Backend / frontend','Java 21 target; Spring Boot 3.5.16; JDBC; React 19 / TypeScript / Vite','Transactions and explicit SQL support reproducible recovery; typed UI data and query caching support operations workflows.'],
['ML / AI','Logistic regression; deterministic read-only tool router','Small evaluated model and inspectable tools avoid granting an autonomous generator mutation authority.'],
['Infrastructure / CI / observability','Compose; Helm; AWS Terraform; GitHub Actions; Prometheus / Grafana / Jaeger','Portable images and manifests; local execution verified. Cloud provisioning and remote CI are still pending.']],[1.28,2.0,3.17],template=3)
heading('5.3 Data Architecture')
figure(er,'Figure 4. Relational core and evidence relationships. Drivers belong to tenants and are optionally assigned to vehicles. Some telemetry relationships are application-enforced to support retention and erasure.')
text('The ownership catalog separates tenant, fleet, driver, manufacturer and vehicle-model attributes. Vehicle state and telemetry summaries are deliberate denormalizations. Raw JSON preserves source shape; revision rows capture canonical values and mapping provenance. Remaining integrity work includes composite tenant foreign keys and PostgreSQL row-level security as defense in depth.')
matrix(['Store / data','Consistency choice','Partition / lifecycle','Verified scope'],[
['PostgreSQL ownership, approvals, raw ledger','Favor consistency; authoritative writes fail if DB unavailable','Tenant/event identity; SQL ledger currently unpartitioned','Real integration and Compose'],
['Kafka raw / normalized / DLT','At-least-once; durability with acks=all; CP-like write rejection under insufficient ISR in RF3 target','tenant:VIN key; raw retention 6h; local RF1','Real broker integration and Compose'],
['Redis / Cassandra projections','Eventual delivery from outbox; stale/missing projections do not authorize changes','Redis latest state TTL 24h; Cassandra tenant/VIN/hour, TTL 24h','Both real integration tests; Redis in core Compose'],
['pgvector / Parquet history','Runbooks tolerate local fallback; exports are snapshot-oriented analytics','Four lexical runbooks; Parquet day/OEM; proposed warm 30d / cold 365d','Vector integration; 18,504-row local export']],[1.43,1.95,1.82,1.25])
text('CAP labels depend on quorum and deployment settings. The local RF1 stack cannot demonstrate partition-tolerant HA. In normal operation, the control plane accepts transaction latency to preserve authorization; disposable projections favor fast eventual access. Cassandra production consistency would be selected explicitly, for example LOCAL_QUORUM versus LOCAL_ONE, after measuring latency and failure behavior.')
text('Capacity worksheet: at 100K events/sec and 1,000 bytes/event, input is 100 MB/sec, 8.64 TB/day and approximately 3.154 PB/year before replication/indexes. A five-minute 300K/sec burst creates 90 million events or 90 GB raw. Consumers holding 100K/sec accumulate 60 GB extra backlog; capacity of 200K/sec after the burst drains it in about ten minutes. At an assumed 4:1 compression, primary raw storage is 2.16 TB/day. Dollar costs require a dated regional quote and approved deployment scope.')
matrix(['Query','Before (ms)','After (ms)','Change Made'],[
['Actual H2 catalog pagination; 100K vehicles','571.64','4.96','Select tenant/VIN page before joins; matching index order; same result'],
['PostgreSQL temporary-table index experiment; 100K rows','3.111','0.063','Composite covering index; same ordered 50 rows; microbenchmark'],
['Dashboard history aggregation','No comparable query-level plan','No comparable query-level plan','V8 transactional rollups; API measurements available. Full before/after EXPLAIN remains pending']],[2.13,.85,.85,2.62],template=4)
text('The template asks for the three slowest production queries with before/after plans. Two local experiments are supplied; the PostgreSQL temporary-table example is not the production fleet query. Staging must capture the actual slowest list, quarantine and dashboard queries under realistic tenant skew and concurrent ingestion. Evidence: sql-h2-query-plans.json and sql-postgres-query-plans.json.')
heading('5.4 Deployment View')
figure(deployment,'Figure 5. Verified local services and proposed deployment. The Kubernetes lane represents supplied assets, not an executed cloud installation.')
text('Compose starts PostgreSQL/pgvector, Kafka, Redis, API and nginx web after local secret generation. Linux containers publish only loopback endpoints. Optional profiles enable Cassandra and monitoring. Helm specifies two API replicas, two web replicas and an HPA maximum of eight, with a pre-existing runtime Secret, resource requests and network policies. Ingress is disabled until an actual host and TLS secret are configured.')
text('AWS Terraform supplies an EKS/KMS/archive foundation and requires existing networking, a chosen Kubernetes version and external data services. A second cloud can reuse OCI images, Helm, SQL migrations and Parquet contracts while replacing infrastructure and service endpoints. Helm lint and Terraform validation passed locally; Kubernetes, cloud failover and second-cloud portability were not executed.')

heading('6. Low-Level Design',page=True)
heading('6.1 Layering & Separation of Concerns')
text('The application uses layered modules with adapters. The normalizer is separate from HTTP and broker transport. FleetService still contains JDBC queries and transaction orchestration, so the prototype does not fully enforce a database-independent application layer. Further repository/port extraction belongs to the distributed evolution, not an unimplemented architecture claim.')
figure(layers,'Figure 6. Current dependency direction. Application services call domain rules and infrastructure adapters; JDBC coupling in FleetService is an explicit prototype trade-off.')
matrix(['Layer','Responsibility','Must Not / Current Boundary'],[
['Presentation / API','DTO validation, JWT/RBAC, HTTP responses, tenant-scoped requests','Must not perform arbitrary normalization. FleetController delegates service operations.'],
['Application / service','Ingestion, approval, replay, privacy and atomic business effects','Currently depends on Spring/JDBC; database-independent ports remain a follow-up.'],
['Domain','VIN rules, field mappings, finite numeric transforms, normalized records','No broker or database calls; shared records use JSON/Jakarta validation types.'],
['Infrastructure','Kafka clients, JDBC, Redis/Cassandra projections, configuration','Sink delivery must not falsely authorize mappings or replace authoritative SQL decisions.']],[1.3,2.35,2.8],template=5)
code('api/       src/main (domain, service, api, config, migrations)\n           src/test (unit, workflow, container, Pact, BDD)\nweb/       src (workspace, profile, themes, governance)\n           tests (appearance, workspace, governance)\nml/        train/evaluation scripts and inference tests\nbatch/     local export and Spark analytics\ninfra/     helm, terraform, observability\nscripts/   launch, evidence, benchmarking, packaging\ntests/     acceptance features and k6 read workload\ndocs/      architecture, ADRs, security, verification\nevidence/  measured JSON reports')
heading('6.2 Design Principles Applied')
text('Single responsibility is clearest in Normalizer, VinValidator, ReplayWorker and PrivacyService. JSON-pointer mapping rules let an engineer extend OEM field interpretation without embedding arbitrary executable expressions. Constructor injection keeps services explicit and testable. FleetService remains broader than ideal; the document does not claim perfect SOLID separation.')
text('Configuration comes from environment variables, logs go to stdout and images are repeatable. Replay/outbox checkpoints live in the database. Mapping caches and the HTTP rate guard are process-local, so full stateless horizontal behavior needs additional work. Idempotency uses a tenant/event key; fail-fast validation rejects impossible values; least privilege is enforced in SecurityConfig and controller methods; shared transformations avoid duplicated OEM parsers.')
heading('6.3 Design Patterns Used')
matrix(['Pattern','Problem It Solves in Your System','Location in Code'],[
['Adapter / mapping strategy','Converts source paths and units into one canonical contract','domain/Normalizer.java; service/EventIntake.java'],
['Transactional outbox','Retains pending sink delivery after the source transaction commits','service/StorageProjection.java; V3 migration'],
['CQRS-style projections','Separates event evidence from latest state and bounded summaries','FleetService.java; TelemetryRollups.java'],
['Circuit breaker / retry','Isolates failing Redis and Cassandra sinks and retries later','ProjectionCircuits.java; StorageProjection.java'],
['Durable checkpointed worker','Resumes replay by persisted page cursor and counters','ReplayWorker.java; V4 migration']],[1.23,2.95,2.27],template=6)
heading('6.4 Interfaces, Contracts & Runtime Flows')
text('The current HTTP contract is implemented by FleetController and TypeScript types/api helpers. No standalone OpenAPI or AsyncAPI specification has been published. Endpoints use /api without a public major-version namespace; OEM schemaVersion and mapping revision are separately explicit. A release should freeze/export the API contract and define a compatibility policy before external adoption.')
matrix(['Interface','Behavior','Access'],[
['GET /api/overview, /vehicles, /alerts','Tenant-scoped operational reads; VIN keyset pagination for vehicles','Authenticated; viewer data masked'],
['POST /api/mappings/{id}/preview and /approve','Sample validation and explicit approval of allowlisted rules','ENGINEER or ADMIN'],
['POST /api/replays; GET /api/replays','Create recovery and inspect durable progress','Mutation ENGINEER/ADMIN'],
['POST /api/ingest','Validate event envelope then intake locally or through Kafka','Authorized engineer/admin'],
['POST /api/assistant','Bounded, read-only evidence tools and audit','Authenticated; role-aware data'],
['POST /api/privacy/erase','Exact repeated VIN confirmation; local purge','ADMIN only']],[2.35,2.75,1.35])
text('Expected business errors use an error message object and appropriate HTTP status. Framework/filter failures may use their own error envelope. Local limits are 30 login requests and 6,000 API requests per source address per minute, with 429 and Retry-After: 60. Declared bodies above 2,000,000 bytes are rejected; proxy/body-stream controls must be completed for production. These are local safeguards, not fleet-scale admission control.')
text('Events are JSON records containing eventId, vin, oem, schemaVersion, eventTime, sequence and payload. Topic keys are tenant:VIN. Topics are fleettruth.raw and fleettruth.normalized, with dead-letter handling for exhausted delivery attempts. Canonical output uses speedKmh, socPct, odometerKm, latitude, longitude, dtc, quality, mappingId and eventTime. Only exact approved OEM/schema versions are interpreted; a changed version enters quarantine. No Avro/Protobuf registry is claimed.')
figure(seq_ingest,'Figure 7. Ingestion and crash/redelivery path. Uniqueness and atomic SQL effects support idempotency across at-least-once Kafka delivery.')
figure(seq_replay,'Figure 8. Unknown-schema failure and human-reviewed recovery. A replay page and its checkpoint commit together; a failed page rolls back.')
heading('6.5 Algorithms & Data Structures')
matrix(['Operation','Algorithm / data structure','Complexity and limits'],[
['VIN and DTC validation','17-position checksum/transliteration; DTC regex','VIN O(17) time/O(1) space; DTC proportional to input; synthetic checksum policy'],
['Signal normalization','JSON pointers and affine scale/offset rules','O(r*d+c) for r rules, pointer depth d and c codes; O(r+c) output'],
['Dedup / latest state','Unique tenant/event key; event-time plus sequence under vehicle lock','Indexed O(log N) identity operation; O(1) state ordering comparison'],
['Replay / pagination','Ordered cursor; 100-event checkpointed pages; VIN keyset','Bounded page memory; indexed traversal. Substring vehicle search can still scan'],
['Drift / retrieval','Five-feature logistic dot product; 128-dimensional exact cosine','O(5) inference after bounded <=200 samples/OEM; O(128*D) for D=4 runbooks']],[1.25,2.5,2.7])
code('INGEST(tenant, event)\n  verify authenticated tenant and owned VIN\n  begin transaction; lock owned vehicle\n  if tenant/eventId exists: return prior outcome\n  retain original event\n  mapping = approved exact OEM/schema contract\n  result = normalize(event, mapping) or quarantine\n  update revision, alert and current state if newer\n  update TOTAL/MINUTE/DAY counters and outbox\n  commit; acknowledge intake\n\nREPLAY(job)\n  lock pending job; load next <=100 retained events\n  normalize using the approved mapping\n  apply only required state/status/counter changes\n  commit page effects with cursor and progress')
text('Measured scale is 100K catalog rows in the SQL experiments and 335,880 stored events at the start of the API sample. No standalone normalizer or VIN throughput microbenchmark was recorded. Graph routing and dynamic programming do not solve this project\'s normalization problem and were not added solely to satisfy example algorithm categories.')

heading('7. Non-Functional Requirements & Performance Benchmarks',page=True)
matrix(['NFR','Target (Case Study)','Achieved','How Measured'],[
['Ingest throughput','100K/sec sustained; 300K/sec burst for 5 min','Unvalidated; demo about 12/sec','Synthetic live subset; no full load/soak run'],
['End-to-end latency','Dashboard <2 sec; critical alert <5 sec','Unvalidated','Two-second UI polling; no instrumented end-to-end percentile'],
['API latency','p95 <200 ms; p99 <500 ms','84.19 / 205.77 ms in bounded sample','90 reads; concurrency 3; 0 errors'],
['Resilience','Recover after broker / pod failure','Replay/idempotency and breaker tests pass; chaos pending','H2/PostgreSQL workflows; breaker state transitions'],
['Availability','99.9%; no single point of failure','Unvalidated; local stack has SPOFs','Single-node Compose; no multi-zone uptime experiment']],[1.05,1.8,1.8,1.8],template=7)
text('Benchmark setup: one Windows development machine, native H2 API, 100K registered vehicles and 335,880 events at the start. Python HTTP reads sampled overview, vehicles and alerts 30 times each, at concurrency three, after one overview warm-up. The simulation continued. The stored report does not record hardware specifications or resource utilization; this limits reproducibility and capacity interpretation.')
figure(ASSETS/'latency.png','Figure 9. Local API latency from evidence/local-api-benchmark.json. Ninety timed reads at concurrency three; nearest-rank percentiles. This sample does not establish a production SLO.')
text('Vehicle-list p95 was 87.32 ms; alerts 11.92 ms; overview 193.57 ms. The earlier mixed p95 of 569.33 ms and p99 of 689.77 ms were observed under different data/cache conditions. They are preserved diagnostic results, not a controlled speedup comparison. The largest architectural limit is per-event relational work plus shared summary-row updates.')
text('The next performance run needs distributed generators, measured consumer lag, unique-event reconciliation, CPU/GC/disk/network metadata, realistic skew and separate sustained, five-minute burst, drain and soak phases. A producer\'s requested rate or broker acknowledgements alone cannot prove downstream zero-loss processing. Stop before exhausting retention or disk capacity.')

heading('8. Security & Compliance',page=True)
text('The local application verifies signed JWTs, enforces roles on API methods and derives HTTP tenant identity from authentication. Viewers cannot inspect raw payloads or mutate mappings; location precision is reduced. Engineers approve/replay; administrators additionally request local erasure with an exact repeated VIN. OIDC authorization-code PKCE and issuer/audience configuration exist but have not been connected to a real organizational IdP.')
matrix(['STRIDE threat','Concrete risk','Current control','Remaining work'],[
['Spoofing','Forged tenant or device','JWT verification; ownership lookup','Device mTLS; trusted producer ACLs'],
['Tampering','Wrong unit or executable mapping','Exact-version rules; finite transforms; sample checks','Signed schema changes; independent approval policy'],
['Repudiation','Approval or access denial is disputed','Approval/replay/API/tool audit records','External append-only sink; internal/denial auditing'],
['Information disclosure','Cross-tenant VIN or precise location','Tenant predicates; viewer masking and raw denial','RLS; encrypted stores; privacy review'],
['Denial of service','Slow sinks, large requests, bursts','Pagination; timeout/retry; independent breakers; read summaries','Distributed quotas; complete payload limits; scale testing'],
['Elevation of privilege','Viewer or assistant approves mapping','Method authorization; no assistant mutation tools','IdP role review; duty separation; artifact signing']],[1.17,1.61,1.9,1.77])
text('Independent Redis and Cassandra circuit breakers use a ten-call window, minimum five calls, 50% failure threshold, 80% slow-call threshold at two seconds, 30-second open interval and two half-open probes. Failed or rejected sink calls keep the outbox pending. This limits repeated sink pressure; it does not replace broker/database HA.')
text('Native HTTP and internal plaintext Compose traffic do not satisfy the production TLS1.3/mTLS/encryption requirements. Production needs TLS endpoints, device credentials, encrypted data services, vault/secret rotation and restricted network paths. Compose secrets are generated locally and excluded from the submission archive. Cassandra\'s local adapter needs authentication/TLS extension before production use.')
text('The erasure API deletes local SQL evidence, registry/state, revisions, alerts and pending outbox rows under vehicle/replay locks, updates summaries and pseudonymizes affected audit resources. Its status is LOCAL_PURGED_EXTERNAL_PENDING. Kafka, Redis, Cassandra, Parquet, backups, traces and operator exports still require coordinated deletion and restore suppression. Hashing is pseudonymization; this prototype does not claim GDPR or DPDP compliance certification.')
text('The assistant uses a deterministic router with bounded tenant-scoped read tools and authored runbooks. Text cannot grant approval or mutate fleet data. An LLM-based prompt-injection surface is not present in this baseline; future LLM integration would need separate untrusted-content boundaries, tool allowlists and adversarial evaluation.')

heading('9. Test Strategy',page=True)
matrix(['Test Type','Tools','No. of Tests','Coverage / Result','In CI?'],[
['Unit / model / query','JUnit / JaCoCo','20 Java','13 normalizer + 5 intelligence + 1 breaker + 1 SQL plan; pass','Defined'],
['Workflow / real integration / contract','Spring, Testcontainers, Pact','40 Java','17 H2 + 18 PostgreSQL + 1 Cassandra + 4 Pact; pass','Defined'],
['Acceptance (BDD)','Cucumber / Spring','3 Java','Three recovery/authorization scenarios pass','Defined'],
['Browser / Python ML','Playwright / pytest','20 per UI target; 2 ML','Native and Docker UI pass; ML holdout/inference pass','Defined; Docker UI run local'],
['Performance / load / soak','HTTP benchmark / SQL; k6 assets','90 timed reads; 2 plan experiments','Bounded measurements only; full load/soak pending','Full load not run'],
['Security','HTTP checks; npm / pip audit / Trivy','13 per API; 2 images','Final app images: 0 high/critical; API 29 medium/low','Scan definitions; no remote run'],
['Privacy / chaos','Workflow and browser tests','Included above','Local erasure/roles checked; external erasure and chaos pending','Partial definition']],[1.27,1.34,.95,2.05,.84],template=8,size=9.5)
text('The Java categories sum to 63 executed tests with zero failures, errors or skips. The integration suite ran actual PostgreSQL, Kafka, Redis, pgvector and Cassandra containers. Pact covers authenticated profile identity and unauthenticated denial with two consumer and two provider tests. Cucumber executes real bindings. Remote CI is defined in .github/workflows/ci.yml but has not been published or run.')
text('Edge cases include duplicate event IDs, out-of-order timestamps, newer quarantined state, invalid VINs/types/ranges, unknown schemas, mapping preview versus approval, replay rollback, summary consistency after erasure, concurrent writers and sink breaker recovery. Browser tests cover approval/replay, profile roles, logout, focus handling, analytics export and both themes at mobile/tablet/desktop widths. These do not constitute universal device or accessibility certification.')
matrix(['Core class','Instruction coverage','Line coverage','Branch coverage'],[
[c['class'].rsplit('.',1)[-1],f"{c['instructionPercent']:.2f}%",f"{c['linePercent']:.2f}%",f"{c['branchPercent']:.2f}%"]
for c in json.loads((ROOT/'evidence/verification.json').read_text())['coverage']
if c['class'] in json.loads((ROOT/'evidence/verification.json').read_text())['criticalCoverageGate']['classes']
],[2.35,1.45,1.25,1.4])
text('The gate enforces 80% line and instruction coverage for these ten explicitly selected core classes. It is not an entire-codebase or branch-coverage claim. Evidence files include verification.json, browser-native-tests.json, browser-docker-tests.json, security-smoke.json, docker-security-smoke.json and the JaCoCo HTML report in api/target/site/jacoco/.')
text('Frontend audit reported zero known findings across 169 dependencies. Python requirements resolved 24 dependencies with zero reported findings; the environment audit was also clean. Trivy inventoried 136 Java packages with zero known findings. The final API image retained 13 medium and 16 low Ubuntu-package findings; the final web image reported zero. Infrastructure images, full SAST and authenticated ZAP remain unverified. All scan results are point-in-time observations.')

heading('10. Observability',page=True)
text('Actuator exposes authorized metrics to Prometheus, Grafana is configured for the local data source, and Micrometer/OpenTelemetry exports HTTP, scheduler and Kafka spans to Jaeger. ECS JSON logs support structured inspection. The captured local-stack report shows a healthy Prometheus scrape, Grafana database health and actual fleettruth traces. Jaeger uses in-memory storage; a durable centralized log index is still pending.')
figure(ROOT/'artifacts/demo/frame-night.png','Figure 10. Working operations dashboard in night mode, showing measured synthetic throughput and contract quality. This is the application dashboard; consumer-lag and infrastructure percentile panels are not shown here.')
text('Troubleshooting a latency spike: identify the affected HTTP endpoint using its histogram; open a corresponding HTTP trace and inspect service/database spans; correlate timestamp and trace ID with structured logs; compare ingest rate, quarantine growth, outbox age and database saturation. If a sink is failing, inspect breaker transitions and pending outbox work before recovery. Verify post-recovery event counts and lag instead of assuming a green health endpoint proves completion.')
text('Current gaps are a Kafka lag exporter, queue-age/stuck-job alerts, end-to-end event freshness histograms, long-term trace retention and centralized log search. Local verification temporarily sampled every trace; default sampling is 10%. Demo Prometheus authentication uses an admin token that expires after four hours and must be renewed through scripts/configure-demo.ps1 -Monitoring.')

heading('11. AI / ML Component (if used)',page=True)
text('The advisory classifier ranks likely schema drift for engineer review. Hard rules remain authoritative for acceptance. Statistical drift scoring complements rules when missing fields, range/type errors and shifts in the battery distribution co-occur, but it cannot determine the true OEM semantics or approve a transformation.')
text('Training uses 8,000 synthetic windows; evaluation uses 2,400 windows with a held-out seed and lower fault severity. Features are missing-field ratio, numeric-type error ratio, range-error ratio, unapproved schema change and relative battery-distribution shift. Parameters are exported for Java inference. Independent seeds and preprocessing fitted on training data reduce direct leakage, but shared synthetic fault families can still make the task easier than unseen OEM incidents.')
matrix(['Metric','Logistic regression','Fixed-threshold baseline'],[
['Positive-class precision','1.000','1.000'],['Positive-class recall','0.9615','0.4590'],
['Positive-class F1','0.9804','0.6292'],['Accuracy','0.9808','0.7308'],
['Holdout support','1,194 drift / 1,206 normal','Same 2,400 windows']],[2.15,2.15,2.15])
text('The holdout confusion matrix contains 1,206 true negatives, zero false positives, 46 false negatives and 1,148 true positives. Results apply to generated fault families only. A pilot needs labeled integration incidents, held-out manufacturers, unfamiliar failures, calibrated probabilities and measured review burden. evidence/ml-evaluation.json contains the complete result.')
text('The assistant routes a bounded question to tenant-scoped fleet, schema, alert, replay or runbook tools; records tool calls; and declines mutation requests. Four authored runbooks use normalized 128-dimensional hashed lexical vectors, with exact cosine in pgvector and a local fallback. It is a deterministic baseline, not an LLM planner or deep semantic embedding system.')
text('There are no external model API charges in this implementation. Hosting and local CPU/memory still have costs; per-request inference/assistant cost and latency were not separately benchmarked. Runtime failure produces an error or documented local runbook fallback. Tool evidence is auditable, but broader language coverage and semantic answer quality need evaluation before extending the assistant.')

heading('12. Architecture Decisions, Risks & Future Enhancements',page=True)
text('ADR 001 - Deterministic contracts. Context: changed units can remain plausible numbers. Options: permissive field matching, autonomous generated mappings or exact approved rules. Decision: quarantine unknown schemas and require human approval after preview. Consequence: data availability temporarily drops while semantic integrity is reviewed. ML remains advisory.')
text('ADR 002 - At-least-once transport with atomic SQL effects. Context: a crash can happen between database and Kafka offset commits. Options: Kafka-only state, cross-store distributed transactions or idempotent relational effects plus an outbox. Decision: unique tenant/event IDs and atomic business effects. Consequence: redelivery is safe for tested effects, but no global exactly-once guarantee spans all stores.')
text('ADR 003 - Consistent control data and eventual projections. Context: approvals require authoritative ownership, while latest-state caches can be rebuilt. Options: one eventual store for everything, one SQL store for all production telemetry, or separated responsibilities. Decision: PostgreSQL control decisions, disposable Redis and time-bucketed Cassandra/lake targets. Consequence: reject authoritative writes during database unavailability; accept bounded staleness in projections. This is the core CAP/PACELC trade-off.')
text('ADR 004 - Deterministic assistant and exact retrieval. Context: only four runbooks and no external model credential are needed for the demo. Options: hosted LLM/general SQL tools, dense vectors/HNSW, or bounded tools plus lexical cosine. Decision: deterministic read-only routing and exact retrieval. Consequence: reproducible evidence and restricted authority, with limited language coverage and possible lexical collisions.')
text('ADR 005 - One deployable prototype with portable boundaries. Context: the hackathon needs an auditable workflow and repeatable local setup. Options: immediate distributed services, provider-specific platform, or modular core with standard protocols. Decision: Spring core, Compose, Helm and an AWS foundation. Consequence: local debugging is manageable, while production worker separation and a second-cloud demonstration remain work.')
text('Known debt includes JDBC-per-event writes, tenant/OEM summary-row contention, process-local rate limits and mapping invalidation, unpartitioned SQL history, incomplete external erasure, missing comprehensive internal-access audit and 29 medium/low API OS findings. Commercial OEM credentials/specifications, real OIDC, device mTLS, managed encryption and HA are not deployed.')
text('Next step 1: run an authorized OEM-specification pilot and collect labeled contract incidents plus operator recovery-time baselines. Next step 2: move bulk telemetry into partitioned stream/time-series/lake storage, separate workers and deploy approved identity/secrets/HA infrastructure. Next step 3: run instrumented sustained/burst/soak and kill-recovery tests, reconcile unique events, complete cross-store erasure and close security findings before broader adoption.')

heading('13. Demo Video (5 Minutes Maximum)',page=True)
text('The supplied local product recording is artifacts/demo/FleetTruth-demo.webm: 114.88 seconds, VP8 WEBM, 1440x1000 at 25 fps, silent. It demonstrates profile/themes, schema drift, sample validation, human approval, replay, alerts, vehicle evidence, assistant, audit and analytics. It fits the duration limit but still needs narration, captions or a final explainer recording and a hosted submission URL. It does not yet meet the template\'s clear-audio and 1080p recommendation.')
matrix(['Time','Segment','What to Show'],[
['0:00-0:30','Problem','Introduce team, mixed OEM formats and the fractional-battery example.'],
['0:30-1:00','Solution','Explain retention, quarantine, approved mapping and recoverable alerts.'],
['1:00-3:00','Live demo','Inject a schema change; inspect source; validate; approve; replay; open recovered evidence.'],
['3:00-4:15','Under the hood','Show architecture, real integration tests, local latency and trace evidence; distinguish capacity targets.'],
['4:15-4:45','Impact and next steps','Explain proposed operator benefit, limitations and the pilot plan.']],[1.05,1.33,4.07],template=9)
text('Hosted video URL: ______________________________________________')
text('The original template limits video to five minutes; a later submission screenshot permits ten. A narrated cut below five minutes satisfies both duration statements. Chapter timing in artifacts/demo/chapters.json is approximate wall-clock timing and can differ slightly from encoded video position.')

heading('14. Repository Checklist')
matrix(['Deliverable','Local status','Location / action'],[
['README and run instructions','Available','README.md; scripts/start-local.ps1; docs/OPERATIONS.md'],
['Compose with simulator','Verified after secret generation','scripts/configure-demo.ps1 then docker compose up --build -d'],
['Code and documentation structure','Available','api, web, ml, batch, infra, scripts, tests, docs, evidence'],
['CI build / test / security definition','Available; remote execution pending','.github/workflows/ci.yml; no published CI URL'],
['Secret hygiene / environment setup','Generated secrets excluded','configure-demo.ps1 creates .env; .env.example not supplied'],
['Repository URL / team commits / tag','Pending','No published repository or v1.0-submission tag'],
['Technical archive','Available','artifacts/submission/FleetTruth-submission.zip; SHA-256 manifest']],[1.9,1.65,2.9])
text('Local demo URLs are http://127.0.0.1:3000 for Docker and http://127.0.0.1:5173 for native mode. The README supplies synthetic local account credentials. These loopback addresses are for reproduction on the host and are not public reviewer links.')

heading('15. Conclusion')
text('FleetTruth demonstrates a complete local workflow for a concrete integration failure: changing vehicle telemetry contracts. It makes source evidence, conversion rules, human approval and recovery results inspectable. The main engineering lessons are that source retention enables repair, replay must preserve ordering/idempotency, and dashboard optimization must stay consistent with mutations such as recovery and erasure.')
text('The strongest evidence is the running native/Docker workflow, actual store integrations, 63 backend passes, 20 browser passes per deployment and scoped coverage/performance reports. The next challenge is to preserve these guarantees while separating bulk telemetry from relational control data and validating the production requirements under real load and failures.')

heading('16. Declarations')
text('All vehicle, driver, manufacturer and telemetry fixtures are synthetic. No private Motorq APIs, production OEM feeds or real vehicle-owner records were used. Motorq is the case-study reference; no affiliation, endorsement or integration partnership is claimed. Header branding is retained from the supplied submission template.')
text('OpenAI Codex assisted with architecture, implementation, debugging, tests, documentation and document preparation. No hosted LLM is called by the runtime assistant. The submitting team is responsible for reviewing the work, supplying its actual identities and representing contributions accurately.')
text('Third-party components are declared through package manifests, lockfiles and the image package inventories. The following major-component license summary should be read with the exact dependency versions and upstream notices. Redis 7.4 and newer Terraform releases use source-available terms; they should not be represented as uniformly permissive open source.')
matrix(['Component family','Declared license family','Location / qualification'],[
['Spring, Kafka, Cassandra, Flyway, Cucumber, OpenTelemetry, Jaeger, Trivy','Apache-2.0','Java manifests and image package inventories; verify component-specific notices'],
['React, Vite, TanStack Query, Recharts; Lucide; Playwright','MIT; ISC; Apache-2.0 respectively','web/package-lock.json and each package LICENSE'],
['PostgreSQL / pgvector','PostgreSQL license','Compose image and extension distributions'],
['H2','MPL-2.0 or EPL-1.0','Embedded native/test database'],
['NumPy / scikit-learn / SciPy; DuckDB / HTTPX; PyArrow','BSD-3-Clause; MIT; Apache-2.0 respectively','requirements.txt and resolved package metadata'],
['Grafana; Redis 7.4; Terraform','AGPL-3.0; RSALv2 or SSPLv1; BSL-1.1 for current releases','Redis/Terraform terms differ from permissive OSS; review redistribution/use obligations']],[2.05,2.1,2.3])

heading('17. Appendix (if any)',page=True)
sub('Evidence index')
matrix(['Artifact','Purpose','Scope'],[
['evidence/verification.json','Backend totals and coverage gate','63 executed tests; ten named core classes'],
['evidence/browser-native-tests.json and browser-docker-tests.json','UI workflow evidence','20 passes per deployment'],
['evidence/local-api-benchmark.json','Client-observed latency sample','90 reads, concurrency 3, nearest-rank percentiles'],
['evidence/sql-h2-query-plans.json and sql-postgres-query-plans.json','Before/after query experiments','100K synthetic rows in each experiment'],
['evidence/ml-evaluation.json','Classifier versus baseline','8K train / 2.4K holdout, synthetic faults'],
['evidence/local-stack.json','Service, metric and trace observations','Local single-node stack'],
['evidence/*audit*.json and *security-smoke.json','Dependency/image findings and HTTP checks','Point-in-time scans; includes before/after image reports'],
['docs/REQUIREMENTS.md; docs/ADRS.md','PDF requirement traceability and full ADRs','Implemented, verified and remaining work identified']],[2.7,1.85,1.9])
sub('Sources')
text('S1. Motorq_Hackathon_Problem_Statement.pdf, supplied 10-page brief, especially pages 6-10. Source of the project constraints, scale targets and deliverables.',size=10)
text('S2. Motorq_Hackathon_Solution_Document_Template.docx, supplied template. All 17 sections and its cover, linked contents and page design are retained in this completed document.',size=10)
text('S3. FleetTruth implementation, tests and measured evidence under api/, web/, ml/, batch/ and evidence/. The verification snapshot was recorded on 2 October 2026 IST.',size=10)
text('S4. Primary technical references already recorded in docs/REFERENCES.md: Apache Kafka design (kafka.apache.org/design/); pgvector documentation (github.com/pgvector/pgvector); Kubernetes HPA (kubernetes.io/docs/concepts/workloads/autoscaling/horizontal-pod-autoscale/); Testcontainers Kafka module (testcontainers.com/modules/kafka/); oidc-client-ts UserManager documentation (authts.github.io/oidc-client-ts/classes/UserManager.html).',size=10)
text('S5. Primary license references checked for the declarations: redis.io/legal/licenses/ (Redis 7.4 RSALv2/SSPLv1); hashicorp.com/en/license-faq (BSL transition); lucide.dev/license (ISC). Other license families are summarized from their package distributions; exact upstream notices govern each release.',size=10)
text('Detailed local source, reports and the demonstration recording are provided as technical artifacts. Raw test XML, browser traces, generated credentials, dependency caches and runtime databases are excluded from the packaged submission. No cloud endpoint, repository URL or remote CI result is asserted.')

doc.core_properties.title='FleetTruth Solution Document'
doc.core_properties.subject='Connected Vehicle Intelligence Hackathon'
doc.core_properties.author=''
doc.core_properties.last_modified_by=''
buffer=BytesIO();doc.save(buffer)
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml'}
destination=OUT/'FleetTruth_Solution.docx'
with ZipFile(REFERENCE) as original, ZipFile(buffer) as generated:
    inventory={info.filename:{'bytes':info.file_size,'sha256':hashlib.sha256(original.read(info.filename)).hexdigest(),
                'editable':info.filename in editable} for info in original.infolist()}
    with ZipFile(destination,'w',ZIP_DEFLATED) as final:
        for info in original.infolist():
            data=generated.read(info.filename) if info.filename in editable else original.read(info.filename)
            final.writestr(info,data)
        for info in generated.infolist():
            if info.filename not in inventory:
                final.writestr(info,generated.read(info.filename))
with ZipFile(destination) as final:
    for name,entry in inventory.items():
        if not entry['editable']:
            assert hashlib.sha256(final.read(name)).hexdigest()==entry['sha256'],name
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest()==EXPECTED_HASH
check=Document(destination)
assert len(check.sections)==1
assert [p.text for p in check.paragraphs if p.style.name.startswith('Heading')]==list(headings)
full_text='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for row in t.rows for c in row.cells)
assert not re.search(r'\[Write your answer|\[Insert diagrams|\[e\.g\.|\[Name|\[Link\]',full_text)
(TMP/'template-inventory.json').write_text(json.dumps(inventory,indent=2)+'\n',encoding='utf-8')
(TMP/'document-build.json').write_text(json.dumps({'sourceSha256':EXPECTED_HASH,'output':str(destination),
    'headings':len(headings),'tables':len(check.tables),'inlineImages':len(check.inline_shapes),
    'words':len(full_text.split()),'preservedOriginalParts':sum(not x['editable'] for x in inventory.values()),
    'pendingCoverFields':['team name','members and roles','repository URL','hosted video URL']},indent=2)+'\n',encoding='utf-8')
print(destination)
print(f'{len(full_text.split())} words, {len(check.tables)} tables, {len(check.inline_shapes)} figures; template fidelity checks passed')
